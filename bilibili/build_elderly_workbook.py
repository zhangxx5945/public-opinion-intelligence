from __future__ import annotations

import csv
import os
import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_ROOT = Path(os.getenv("CRAWLER_OUTPUT_ROOT", PROJECT_ROOT / "outputs")) / "bilibili"
SOURCE_DIR = Path(os.getenv("BILIBILI_ELDERLY_SOURCE_DIR", OUTPUT_ROOT / "raw" / "elderly_2025"))
OUTPUT_PATH = Path(os.getenv("BILIBILI_ELDERLY_WORKBOOK", OUTPUT_ROOT / "bilibili_elderly_2025.xlsx"))
COMPAT_PATH = Path(os.getenv("BILIBILI_ELDERLY_COMPAT_WORKBOOK", OUTPUT_ROOT / "bilibili_elderly_2025_compat.xlsx"))

COLUMNS = [
    ("关键词", "sheet_keyword"),
    ("视频ID", "video_id"),
    ("标题", "title"),
    ("作者", "author"),
    ("发布时间", "published_at"),
    ("时长", "duration"),
    ("播放量", "views"),
    ("点赞数", "likes"),
    ("评论数", "comments"),
    ("视频链接", "url"),
    ("抓取时间", "crawled_at"),
]

INTEGER_FIELDS = {"views", "likes", "comments"}
DATETIME_FIELDS = {"published_at", "crawled_at"}
WIDTHS = [26, 17, 52, 24, 20, 15, 12, 11, 11, 44, 20]


def filename_keyword(path: Path) -> str:
    return re.sub(r"^\d+_", "", path.stem)


def convert(field: str, value: str):
    value = (value or "").strip()
    if not value:
        return None
    if field in INTEGER_FIELDS:
        try:
            return int(float(value))
        except ValueError:
            return value
    if field in DATETIME_FIELDS:
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00")).replace(tzinfo=None)
        except ValueError:
            return value
    return value


def collect_rows() -> list[list[object]]:
    csv_files = sorted(SOURCE_DIR.glob("*.csv"))
    if len(csv_files) != 29:
        raise RuntimeError(f"预期29个原始CSV，实际找到{len(csv_files)}个")

    rows: list[list[object]] = []
    for path in csv_files:
        fallback_keyword = filename_keyword(path)
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            for record in csv.DictReader(handle):
                keyword = (record.get("sheet_keyword") or fallback_keyword).strip()
                values = []
                for _, field in COLUMNS:
                    raw_value = keyword if field == "sheet_keyword" else record.get(field, "")
                    values.append(convert(field, raw_value))
                rows.append(values)
    return rows


def build_workbook(rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "全部数据"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"

    sheet.append([label for label, _ in COLUMNS])
    for row in rows:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    bottom_border = Border(bottom=Side(style="thin", color="B4C6E7"))
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = bottom_border
    sheet.row_dimensions[1].height = 28

    for index, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    max_row = len(rows) + 1
    sheet.auto_filter.ref = f"A1:K{max_row}"
    for row in sheet.iter_rows(min_row=2, max_row=max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column == 3)

    for column in (5, 11):
        for cells in sheet.iter_cols(min_col=column, max_col=column, min_row=2, max_row=max_row):
            cells[0].number_format = "yyyy-mm-dd hh:mm"
    for column in (7, 8, 9):
        for cells in sheet.iter_cols(min_col=column, max_col=column, min_row=2, max_row=max_row):
            cells[0].number_format = "#,##0"
    for row_index in range(2, max_row + 1):
        link_cell = sheet.cell(row=row_index, column=10)
        if link_cell.value:
            link_cell.hyperlink = str(link_cell.value)
            link_cell.style = "Hyperlink"

    workbook.save(OUTPUT_PATH)
    shutil.copyfile(OUTPUT_PATH, COMPAT_PATH)


def verify(expected_rows: int) -> None:
    for path in (OUTPUT_PATH, COMPAT_PATH):
        workbook = load_workbook(path, read_only=False, data_only=False)
        if workbook.sheetnames != ["全部数据"]:
            raise RuntimeError(f"{path.name} 工作表校验失败：{workbook.sheetnames}")
        sheet = workbook["全部数据"]
        if sheet.max_row - 1 != expected_rows or sheet.max_column != len(COLUMNS):
            raise RuntimeError(f"{path.name} 数据范围校验失败")
        headers = [sheet.cell(1, column).value for column in range(1, len(COLUMNS) + 1)]
        if headers != [label for label, _ in COLUMNS]:
            raise RuntimeError(f"{path.name} 表头校验失败")
        workbook.close()


if __name__ == "__main__":
    combined_rows = collect_rows()
    build_workbook(combined_rows)
    verify(len(combined_rows))
    unique_ids = len({row[1] for row in combined_rows if row[1]})
    print({
        "worksheets": 1,
        "rows": len(combined_rows),
        "unique_video_ids": unique_ids,
        "columns": len(COLUMNS),
        "output": str(OUTPUT_PATH),
        "compat_output": str(COMPAT_PATH),
    })
